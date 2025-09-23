#!/usr/bin/env python3
"""
Generic Markdown→XLSX converter for simple pipe tables.
- Creates one sheet per top-level section (## Heading)
- If multiple tables exist under the same section, appends them vertically with a blank row

Usage:
  ./helper/md_to_xlsx.py /abs/path/to/input.md /abs/path/to/output.xlsx

Requires: openpyxl
"""
import sys
import re
from pathlib import Path
from typing import List, Tuple
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def auto_fit(ws):
    widths = {}
    for row in ws.iter_rows(values_only=True):
        for i, v in enumerate(row, start=1):
            l = len(str(v)) if v is not None else 0
            widths[i] = max(widths.get(i, 0), l)
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), 80)


def parse_sections(md: str) -> List[Tuple[str, str]]:
    # Split by H2 (## ) sections; keep title and body
    parts = re.split(r"^##\s+(.+)$", md, flags=re.MULTILINE)
    sections = []
    if parts[0].strip():
        sections.append(("Overview", parts[0]))
    for i in range(1, len(parts), 2):
        title = parts[i].strip()
        body = parts[i + 1]
        sections.append((title, body))
    return sections


def extract_tables(md: str) -> List[List[List[str]]]:
    lines = md.splitlines()
    tables = []
    i = 0
    while i < len(lines):
        if "|" in lines[i]:
            if i + 1 < len(lines) and re.match(r"^\s*\|?\s*[-: ]+\|[-:| ]+\s*$", lines[i + 1]):
                # gather
                tlines = [lines[i]]
                i += 2
                while i < len(lines) and "|" in lines[i] and not re.match(r"^\s*$", lines[i]):
                    tlines.append(lines[i])
                    i += 1
                # parse
                rows = []
                for tl in tlines:
                    if re.match(r"^\s*\|?\s*[-: ]+\|[-:| ]+\s*$", tl):
                        continue
                    parts = [p.strip() for p in tl.strip().strip("|").split("|")]
                    rows.append(parts)
                if rows:
                    tables.append(rows)
                continue
        i += 1
    return tables


def to_sheet_name(title: str) -> str:
    name = re.sub(r"[\\/*?:\[\]]", " ", title).strip()
    return (name[:28] or "Sheet").rstrip()


def convert(md_path: Path, xlsx_path: Path):
    md = md_path.read_text(encoding="utf-8")
    sections = parse_sections(md)
    wb = Workbook()
    wb.remove(wb.active)

    for title, body in sections:
        tables = extract_tables(body)
        if not tables:
            # skip empty sections
            continue
        sheet_name = to_sheet_name(title)
        base = sheet_name
        idx = 1
        while sheet_name in wb.sheetnames:
            idx += 1
            sheet_name = f"{base[:25]} {idx}"
        ws = wb.create_sheet(sheet_name)
        header_font = Font(bold=True)
        center = Alignment(horizontal="center")
        row_cursor = 1
        for t in tables:
            for r_idx, row in enumerate(t, start=0):
                for c_idx, val in enumerate(row, start=1):
                    ws.cell(row=row_cursor, column=c_idx, value=val)
                if r_idx == 0:
                    # header row styling
                    for c_idx in range(1, len(row) + 1):
                        cell = ws.cell(row=row_cursor, column=c_idx)
                        cell.font = header_font
                        cell.alignment = center
                row_cursor += 1
            # blank line between tables
            row_cursor += 1
        auto_fit(ws)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(xlsx_path))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: md_to_xlsx.py input.md output.xlsx", file=sys.stderr)
        sys.exit(2)
    convert(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
    print(f"Written: {Path(sys.argv[2]).resolve()}")
